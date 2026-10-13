"""Builds the copy of "[Foundever Copy] Disconnected Calls - All Markets" with the
client's reporting requirements (xlsx to import into Google Sheets).

    python3 build_copy.py <export.xlsx> <output.xlsx> [--lo] [--test-inputs <json>]

Input: the xlsx export of the Google Sheet (File > Download > xlsx).
Output: the same workbook with
  - review columns U:AA on the five language tabs (root cause, callback, action date,
    review check, repeat count), with "Callback made?" pre-filled from the old review logs;
  - new tabs: Read Me, Daily Report, Daily Summary, Endpoint Check, All Cases,
    Repeat Offenders, Process Tracker, Settings;
  - new dropdown lists on the hidden Validation tab and two new Action Taken values.

The workbook is edited at XML level: every other part of the export (comments, tasks,
images, filter views, tables, Google formulas wrapped in __xludf.DUMMYFUNCTION) is copied
byte for byte. Formulas use Google Sheets syntax (ARRAYFORMULA, LET, FILTER, VSTACK...):
open the result with Google Sheets.

--lo writes the same formulas in Excel/LibreOffice syntax (array formulas, _xlfn. prefixes)
so that LibreOffice can recalculate the file for testing. --test-inputs fills test values
in input cells (test file only, never for the real copy).
"""
import datetime as dt
import html
import json
import re
import sys
import zipfile
from collections import defaultdict

import openpyxl

args = sys.argv[1:]
SRC, OUT = args[0], args[1]
DIALECT = "lo" if "--lo" in args else "google"
TEST_INPUTS = json.load(open(args[args.index("--test-inputs") + 1])) if "--test-inputs" in args else None

FIRST_DAY = dt.date(2026, 9, 7)   # first Call Day in RAW
LAST_DAY = dt.date(2027, 6, 30)   # last day of the Daily Summary / Endpoint Check grids
SITE_LAST_ROW = 2728              # language tabs are formatted (tables, lists) down to this row
RAW_LAST_ROW = 6655               # RAW!P:R lookups run down to this row
CASES_LAST_ROW = 14000            # All Cases: room for the five tabs (5 x 2727 rows)
RO_ROWS = 1000                    # Repeat Offenders rows
PT_ROWS = 500                     # Process Tracker rows

SITES = [  # tab, sheet part, team (template "Cost Center"), cost center, RAW cost center contains, RAW language
    ("São Paulo", "xl/worksheets/sheet11.xml", "São Paulo", "Sitel SaoPaulo", "Paulo", "Portuguese"),
    ("Portugal PT", "xl/worksheets/sheet9.xml", "Portugal", "Sitel Portugal", "Portugal", "Portuguese"),
    ("Portugal SP", "xl/worksheets/sheet13.xml", "Portugal", "Sitel Portugal", "Portugal", "Spanish"),
    ("Seville", "xl/worksheets/sheet14.xml", "Seville", "Sitel Spain", "Spain", "Spanish"),
    ("Milan", "xl/worksheets/sheet16.xml", "Milan", "Sitel Milan", "Milan", "Italian"),
]
OTHER = "Other (no tab)"
VALIDATION_XML = "xl/worksheets/sheet18.xml"

LISTS = {  # Validation tab column -> (header, values)
    "D": ("Root cause", [
        "Technical - Headset", "Technical - Microphone", "Technical - Unmute / mute",
        "Technical - Transfer", "Technical - Network / connection", "Technical - Genesys / pop-up error",
        "Technical - PC / browser", "Technical - Other",
        "Behaviour - Closed or minimised pop-up", "Behaviour - Hung up on customer",
        "Behaviour - Call avoidance", "Behaviour - Other",
        "Valid - Customer ended the call", "Valid - Natural end of call", "Under investigation"]),
    "E": ("Callback made?", ["Yes", "No", "N/A"]),
    "F": ("No-callback reason", [
        "Customer called back first", "Continued by message", "Case not assigned to agent",
        "Agent did not attempt callback", "Other"]),
    "G": ("Yes / No", ["Yes", "No"]),
    "H": ("Process stage", [
        "Monitoring", "Coaching 1", "Coaching 2", "Written warning", "Final warning",
        "Removal from campaign", "HR process", "Closed"]),
    "I": ("Endpoint check cause", [
        "Network / ISP", "VPN / remote access", "Headset / USB device", "PC / browser performance",
        "Genesys / telephony incident", "Pop-up closed or minimised by agents", "Other"]),
}
NEW_ACTIONS = ["IT ticket / technical fix", "Feedback / reminder sent"]  # added to Action Taken (R)
OLD_ACTIONS = "Coached,Removal,N/A,Pending"
ROOT = LISTS["D"][1]
N_TECH = sum(v.startswith("Technical") for v in ROOT)
N_BEH = sum(v.startswith("Behaviour") for v in ROOT)
TECH_RANGE = f"Validation!$D$2:$D${1 + N_TECH}"
BEH_RANGE = f"Validation!$D${2 + N_TECH}:$D${1 + N_TECH + N_BEH}"


def lst(col):
    return f"Validation!${col}$2:${col}${1 + len(LISTS[col][1])}"


# ----------------------------------------------------------------------------- helpers
def col_num(c):
    n = 0
    for ch in c:
        n = n * 26 + ord(ch) - 64
    return n


def col_let(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def esc(s):
    return html.escape(str(s), quote=False)


def serial(d):
    return (d - dt.date(1899, 12, 30)).days


MONTHS = '"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"'


def en_day(e, short=False):
    """Weekday name in English whatever the spreadsheet locale (TEXT(...,"dddd") follows the locale)."""
    if short:
        return f'CHOOSE(WEEKDAY({e}),"Sun","Mon","Tue","Wed","Thu","Fri","Sat")'
    return f'CHOOSE(WEEKDAY({e}),"Sunday","Monday","Tuesday","Wednesday","Thursday","Friday","Saturday")'


def en_date(e):
    """'11 Oct 2026' in English whatever the spreadsheet locale."""
    return f'DAY({e})&" "&CHOOSE(MONTH({e}),{MONTHS})&" "&YEAR({e})'


class F:
    """A formula in Google Sheets syntax; `array` = range of an array result (or the cell)."""

    def __init__(self, text, array=None):
        self.text = text.lstrip("=")
        self.array = array
        depth = 0
        for part in re.split(r'"[^"]*"', self.text):  # outside string literals
            for ch in part:
                depth += {"(": 1, ")": -1}.get(ch, 0)
                assert depth >= 0, self.text
        assert depth == 0 and self.text.count('"') % 2 == 0, self.text


XLFN = ["UNIQUE", "LET", "XLOOKUP", "TEXTJOIN", "VSTACK", "HSTACK", "MAXIFS", "MINIFS", "IFS"]


def lo_text(t):
    t = re.sub(r"(?<![\w.])(FILTER|SORT)\(", r"_xlfn._xlws.\1(", t)
    t = re.sub(r"(?<![\w.])(" + "|".join(XLFN) + r")\(", r"_xlfn.\1(", t)
    return re.sub(r"(?<![\w.])(v_\w+)", r"_xlpm.\1", t)


def render_formula(f):
    """-> (formula text, extra attributes of <f>)"""
    t = f.text.replace("{DESC}", "FALSE" if DIALECT == "google" else "-1")
    t = t.replace("{ASC}", "TRUE" if DIALECT == "google" else "1")
    if DIALECT == "google":
        return (f"ARRAYFORMULA({t})" if f.array else t), ""
    t = lo_text(t)
    return t, (f' t="array" ref="{f.array}"' if f.array else "")


def cell_xml(ref, value, style):
    s = f' s="{style}"' if style is not None else ""
    if value is None or value == "":
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, F):
        text, attrs = render_formula(value)
        return f'<c r="{ref}"{s}><f{attrs}>{esc(text)}</f></c>'
    if isinstance(value, bool):
        return f'<c r="{ref}"{s} t="b"><v>{int(value)}</v></c>'
    if isinstance(value, (dt.date, dt.datetime)):
        return f'<c r="{ref}"{s}><v>{serial(value)}</v></c>'
    if isinstance(value, (int, float)):
        return f'<c r="{ref}"{s}><v>{value}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{esc(value)}</t></is></c>'


# ----------------------------------------------------------------------------- styles
class Styles:
    def __init__(self, xml):
        self.xml = xml
        self.new = {k: [] for k in ("numFmts", "fonts", "fills", "borders", "cellXfs", "dxfs")}
        self.count = {}
        for k in self.new:
            m = re.search(rf"<{k} count=\"(\d+)\"", xml)
            self.count[k] = int(m.group(1)) if m else 0
        self.cache = {}

    def _add(self, kind, x):
        key = (kind, x)
        if key not in self.cache:
            self.new[kind].append(x)
            self.cache[key] = self.count[kind] + len(self.new[kind]) - 1
        return self.cache[key]

    def numfmt(self, code):
        key = ("numFmtCode", code)
        if key not in self.cache:
            ids = [int(i) for i in re.findall(r'numFmtId="(\d+)"', re.search(r"<numFmts.*?</numFmts>", self.xml, re.S).group(0))]
            nid = max(ids + [163]) + 1 + len(self.new["numFmts"])
            self.new["numFmts"].append(f'<numFmt numFmtId="{nid}" formatCode="{html.escape(code)}"/>')
            self.cache[key] = nid
        return self.cache[key]

    def font(self, size=10, bold=False, italic=False, color="FF000000"):
        return self._add("fonts", f'<font>{"<b/>" if bold else ""}{"<i/>" if italic else ""}<sz val="{size}"/>'
                                  f'<color rgb="{color}"/><name val="Calibri"/><scheme val="minor"/></font>')

    def fill(self, rgb):
        if rgb is None:
            return 0
        return self._add("fills", f'<fill><patternFill patternType="solid"><fgColor rgb="{rgb}"/>'
                                  f'<bgColor rgb="{rgb}"/></patternFill></fill>')

    def border(self, rgb):
        if rgb is None:
            return 0
        side = f'style="thin"><color rgb="{rgb}"/>'
        return self._add("borders", f"<border><left {side}</left><right {side}</right><top {side}</top>"
                                    f"<bottom {side}</bottom><diagonal/></border>")

    def xf(self, size=10, bold=False, italic=False, color="FF000000", fill=None, border="FFD9D9D9",
           num=0, h=None, v="center", wrap=False):
        al = f'<alignment{f" horizontal=\"{h}\"" if h else ""} vertical="{v}"{" wrapText=\"1\"" if wrap else ""}/>'
        return self._add("cellXfs", f'<xf borderId="{self.border(border)}" fillId="{self.fill(fill)}" '
                                    f'fontId="{self.font(size, bold, italic, color)}" numFmtId="{num}" xfId="0" '
                                    f'applyAlignment="1" applyBorder="1" applyFill="1" applyFont="1" '
                                    f'applyNumberFormat="1">{al}</xf>')

    def dxf(self, font=None, fill=None, bold=False):
        f = f'<font>{"<b/>" if bold else ""}{f"<color rgb=\"{font}\"/>" if font else ""}</font>'
        fl = (f'<fill><patternFill patternType="solid"><fgColor rgb="{fill}"/><bgColor rgb="{fill}"/>'
              f'</patternFill></fill>') if fill else ""
        return self._add("dxfs", f"<dxf>{f}{fl}<border/></dxf>")

    def render(self):
        x = self.xml
        for k, items in self.new.items():
            if not items:
                continue
            total = self.count[k] + len(items)
            x = re.sub(rf'<{k} count="\d+"', f'<{k} count="{total}"', x, count=1)
            x = x.replace(f"</{k}>", "".join(items) + f"</{k}>", 1)
        return x


# ----------------------------------------------------------------------------- input
zin = zipfile.ZipFile(SRC)
parts = {n: zin.read(n) for n in zin.namelist()}
ST = Styles(parts["xl/styles.xml"].decode("utf-8"))

DATE_FMT = 172  # d"-"mmm"-"yyyy, already used by Validation Date
PCT = ST.numfmt("0.0%")
INT = 1

TEAL, TEAL_LIGHT = "FF26A69A", "FFE0F2F1"   # input columns (same teal as the review tables)
SLATE, GREY_LIGHT = "FF78909C", "FFF3F3F3"  # formula columns (same blue-grey as the case tables)
X = {
    "title": ST.xf(size=14, bold=True, color="FF37474F", border=None),
    "sub": ST.xf(size=10, italic=True, color="FF666666", border=None),
    "note": ST.xf(size=10, color="FF434343", border=None, v="top", wrap=True),
    "section": ST.xf(size=11, bold=True, color="FF37474F", border=None),
    "group": ST.xf(size=10, bold=True, color="FF37474F", fill="FFCFD8DC", h="center"),
    "h_calc": ST.xf(bold=True, color="FFFFFFFF", fill=SLATE, h="center", wrap=True),
    "h_input": ST.xf(bold=True, color="FFFFFFFF", fill=TEAL, h="center", wrap=True),
    "calc": ST.xf(fill=GREY_LIGHT),
    "calc_c": ST.xf(fill=GREY_LIGHT, h="center"),
    "calc_int": ST.xf(fill=GREY_LIGHT, num=INT, h="center"),
    "calc_pct": ST.xf(fill=GREY_LIGHT, num=PCT, h="center"),
    "calc_date": ST.xf(fill=GREY_LIGHT, num=DATE_FMT, h="center"),
    "calc_wrap": ST.xf(fill=GREY_LIGHT, wrap=True, v="top"),
    "calc_bold": ST.xf(fill=GREY_LIGHT, bold=True),
    "key": ST.xf(color="FFB7B7B7", border=None),
    "input": ST.xf(fill=TEAL_LIGHT),
    "input_c": ST.xf(fill=TEAL_LIGHT, h="center"),
    "input_int": ST.xf(fill=TEAL_LIGHT, num=INT, h="center"),
    "input_pct": ST.xf(fill=TEAL_LIGHT, num=PCT, h="center"),
    "input_date": ST.xf(fill=TEAL_LIGHT, num=DATE_FMT, h="center"),
    "input_wrap": ST.xf(fill=TEAL_LIGHT, wrap=True, v="top"),
    "text": ST.xf(),
    "text_b": ST.xf(bold=True),
    "text_i": ST.xf(italic=True, color="FF666666"),
    "date": ST.xf(num=DATE_FMT, h="center"),
    "plain": ST.xf(border=None),
    "wrap": ST.xf(wrap=True, v="top"),
    "end": ST.xf(italic=True, color="FFBFBFBF", border=None),
    # language tabs: 11 pt like the existing case rows
    "s_calc": ST.xf(size=11, fill=GREY_LIGHT),
    "s_calc_c": ST.xf(size=11, fill=GREY_LIGHT, h="center"),
    "s_input": ST.xf(size=11),
    "s_input_date": ST.xf(size=11, num=DATE_FMT),
    "s_prefill": ST.xf(size=11, italic=True, color="FF1F4E79"),
    "s_h_calc": ST.xf(bold=True, color="FFFFFFFF", fill=SLATE, h="center", wrap=True),
    "s_h_input": ST.xf(bold=True, color="FFFFFFFF", fill=TEAL, h="center", wrap=True),
}
DXF = {
    "green": ST.dxf(font="FF274E13", fill="FFD9EAD3"),
    "red": ST.dxf(font="FF990000", fill="FFF4CCCC"),
    "orange": ST.dxf(font="FF783F04", fill="FFFCE5CD"),
    "yellow": ST.dxf(font="FF7F6000", fill="FFFFF2CC"),
    "grey": ST.dxf(font="FF999999"),
    "zero": ST.dxf(font="FFD9D9D9"),
    "weekend": ST.dxf(fill="FFECEFF1"),
    "boldred": ST.dxf(font="FFCC0000", bold=True),
}

wb_vals = openpyxl.load_workbook(SRC, read_only=True, data_only=True)


def site_ids(tab):
    """Id Interaction -> row on a language tab (cached values of the QUERY spill)."""
    out = {}
    for i, r in enumerate(wb_vals[tab].iter_rows(min_row=2, max_row=SITE_LAST_ROW, max_col=1, values_only=True), 2):
        if r[0]:
            out[str(r[0]).strip()] = i
    return out


IDS = {s[0]: site_ids(s[0]) for s in SITES}


# ----------------------------------------------------------------------------- pre-fill from old review logs
def norm_cb(v):
    v = str(v or "").strip().lower()
    if v in ("yes", "y"):
        return "Yes"
    if v in ("no", "n"):
        return "No"
    if v.startswith("n/a") or v == "na":
        return "N/A"
    return None


def legacy_rows(tab, id_cols, val_col):
    for r in wb_vals[tab].iter_rows(min_row=2, values_only=True):
        for c in id_cols:
            v = r[c] if c < len(r) else None
            if isinstance(v, str) and len(v.strip()) > 30:
                yield v.strip(), r
                break


prefill = defaultdict(dict)  # id -> {col: value}
conflict = set()


def put(i, col, val):
    if val is None or i in conflict:
        return
    old = prefill[i].get(col)
    if old is not None and old != val:
        conflict.add(i)
        prefill.pop(i, None)
        return
    prefill[i][col] = val


for tab, idc, cb in (("PT_São Paulo", (1, 11), 18), ("PT_Portugal", (1, 11), 18), ("SP_Seville & Portugal", (1,), 22)):
    for i, r in legacy_rows(tab, idc, cb):
        put(i, "W", norm_cb(r[cb] if cb < len(r) else None))
MILAN_CAUSE = {"internet issue": "Technical - Network / connection",
               "intentional closed - avoidance": "Behaviour - Call avoidance"}
MILAN_WHY = {"user called back within 2 minutes": "Customer called back first",
             "continued via obm": "Continued by message", "avoidance": "Agent did not attempt callback"}
for i, r in legacy_rows("IT_Milan", (1,), 14):
    cb = norm_cb(r[14])
    put(i, "W", cb)
    put(i, "V", MILAN_CAUSE.get(str(r[13] or "").strip().lower()))
    if cb == "No":
        put(i, "X", MILAN_WHY.get(str(r[15] or "").strip().lower()))
PREFILL = {s[0]: {} for s in SITES}  # tab -> row -> {col: value}
for tab, ids in IDS.items():
    for i, row in ids.items():
        if i in prefill and prefill[i]:
            PREFILL[tab][row] = dict(prefill[i])
LEGACY_ROWS = {t: {r: dict(v) for r, v in rows.items()} for t, rows in PREFILL.items()}
if TEST_INPUTS:
    for tab, rows in TEST_INPUTS.get("site", {}).items():
        for row, vals in rows.items():
            PREFILL[tab].setdefault(int(row), {}).update(vals)
PREFILL_COUNT = {t: sum("W" in v for v in rows.values()) for t, rows in PREFILL.items()}


# ----------------------------------------------------------------------------- language tabs
def sheet_cf(rules, start_priority):
    out, p = [], start_priority
    for sqref, kind, formula, dxf in rules:
        if kind == "expr":
            out.append(f'<conditionalFormatting sqref="{sqref}"><cfRule type="expression" dxfId="{dxf}" '
                       f'priority="{p}"><formula>{esc(formula)}</formula></cfRule></conditionalFormatting>')
        else:  # (operator, value)
            op, val = formula
            out.append(f'<conditionalFormatting sqref="{sqref}"><cfRule type="cellIs" dxfId="{dxf}" '
                       f'priority="{p}" operator="{op}"><formula>{esc(val)}</formula></cfRule></conditionalFormatting>')
        p += 1
    return "".join(out)


def dv_xml(sqref, kind, f1, strict=True, op=None):
    a = f'type="{kind}" allowBlank="1"' + (' showErrorMessage="1"' if strict else "") + (f' operator="{op}"' if op else "")
    return f'<dataValidation {a} sqref="{sqref}"><formula1>{esc(f1)}</formula1></dataValidation>'


SITE_COLS = [  # col, header, kind, width
    ("U", "Case type", "calc", 17),
    ("V", "Root cause", "input", 30),
    ("W", "Callback made?", "input", 11),
    ("X", "No-callback reason", "input", 26),
    ("Y", "Action date", "input_date", 12),
    ("Z", "Review check", "calc", 30),
    ("AA", "Agent cases (repeat window)", "calc_c", 12),
]
L = SITE_LAST_ROW
SITE_FORMULAS = {
    "U": F(f'IF(A2:A{L}="","",IF(G2:G{L}=1,"Hung up (final)",IF(F2:F{L}="endpoint","Pop-up disconnect",'
           f'"Other disconnect")))', f"U2:U{L}"),
    "Z": F(f'IF(A2:A{L}="","",IF(P2:P{L}="","Not reviewed",IF(V2:V{L}="","Root cause missing",'
           f'IF(((P2:P{L}="Technical issues")*(LEFT(V2:V{L},9)<>"Technical"))'
           f'+((P2:P{L}="Specialist issues")*(LEFT(V2:V{L},9)<>"Behaviour"))'
           f'+((LEFT(P2:P{L},10)="False Flag")*(LEFT(V2:V{L},5)<>"Valid"))'
           f'+((P2:P{L}="Under Investigation")*(V2:V{L}<>"Under investigation"))>0,'
           f'"Root cause differs from Validation",IF(W2:W{L}="","Callback missing",'
           f'IF((W2:W{L}="No")*(X2:X{L}=""),"No-callback reason missing",IF(R2:R{L}="","Action missing",'
           f'IF((R2:R{L}<>"N/A")*(R2:R{L}<>"Pending")*(Y2:Y{L}=""),"Action date missing","Complete"))))))))',
           f"Z2:Z{L}"),
    "AA": F(f'IF(A2:A{L}="","",COUNTIFS(I2:I{L},I2:I{L},C2:C{L},"<="&C2:C{L},'
            f'C2:C{L},">"&(C2:C{L}-Settings!$B$6)))', f"AA2:AA{L}"),
}


def patch_site(tab, xml):
    st_h = {"calc": X["s_h_calc"], "calc_c": X["s_h_calc"], "input": X["s_h_input"], "input_date": X["s_h_input"]}
    st_b = {"calc": X["s_calc"], "calc_c": X["s_calc_c"], "input": X["s_input"], "input_date": X["s_input_date"]}
    fills = PREFILL[tab]

    def row_cells(r):
        out = []
        for col, head, kind, _ in SITE_COLS:
            ref = f"{col}{r}"
            if r == 1:
                out.append(cell_xml(ref, head, st_h[kind]))
            elif r == 2 and col in SITE_FORMULAS:
                out.append(cell_xml(ref, SITE_FORMULAS[col], st_b[kind]))
            else:
                v = fills.get(r, {}).get(col) if kind.startswith("input") else None
                if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                    v = dt.date.fromisoformat(v)
                legacy = r in LEGACY_ROWS[tab] and col in LEGACY_ROWS[tab][r] and v == LEGACY_ROWS[tab][r][col]
                out.append(cell_xml(ref, v, X["s_prefill"] if legacy else st_b[kind]))
        return "".join(out)

    def repl(m):
        r = int(m.group(1))
        if r > SITE_LAST_ROW:
            return m.group(0)
        body = m.group(3)
        if "<c r=\"U" in body:
            raise SystemExit(f"{tab}: row {r} already has cells after T")
        return f'<row r="{r}"{m.group(2)}>{body}{row_cells(r)}</row>'

    assert xml.count("QUERY(RAW!A1:R2728,") == 1, tab
    xml = xml.replace("QUERY(RAW!A1:R2728,", f"QUERY(RAW!A1:R{RAW_LAST_ROW},")
    xml, n = re.subn(r'<row r="(\d+)"([^>]*)>(.*?)</row>', repl, xml, flags=re.S)
    assert n >= SITE_LAST_ROW, (tab, n)
    cols = "".join(f'<col customWidth="1" min="{col_num(c)}" max="{col_num(c)}" width="{w}"/>' for c, _, _, w in SITE_COLS)
    xml = xml.replace("</cols>", cols + "</cols>", 1)
    # Action Taken: two more values (same list on every row range)
    new_list = OLD_ACTIONS + "," + ",".join(NEW_ACTIONS)
    assert xml.count(f"&quot;{OLD_ACTIONS}&quot;") >= 1 or tab == "Milan", tab
    xml = xml.replace(f"&quot;{OLD_ACTIONS}&quot;", f"&quot;{new_list}&quot;")
    dvs = (dv_xml(f"V2:V{L}", "list", lst("D")) + dv_xml(f"W2:W{L}", "list", lst("E"))
           + dv_xml(f"X2:X{L}", "list", lst("F")) + dv_xml(f"Y2:Y{L}", "date", "1", op="greaterThan"))
    xml = xml.replace("</dataValidations>", dvs + "</dataValidations>", 1)
    cf = sheet_cf([
        (f"Z2:Z{L}", "expr", '$Z2="Complete"', DXF["green"]),
        (f"Z2:Z{L}", "expr", '$Z2="Not reviewed"', DXF["grey"]),
        (f"Z2:Z{L}", "expr", 'AND($Z2<>"",$Z2<>"Complete",$Z2<>"Not reviewed")', DXF["orange"]),
        (f"AA2:AA{L}", "expr", 'AND(ISNUMBER($AA2),$AA2>=INDIRECT("Settings!B7"))', DXF["boldred"]),
    ], 20)
    i = xml.rfind("</conditionalFormatting>")
    i = i + len("</conditionalFormatting>") if i > 0 else xml.find("<dataValidations")
    xml = xml[:i] + cf + xml[i:]
    if tab == "Portugal SP":  # sheet filter also covers the new columns
        xml = xml.replace('<autoFilter ref="$A$1:$T$2728"/>', '<autoFilter ref="$A$1:$AA$2728"/>', 1)
    return xml


for tab, part, *_ in SITES:
    parts[part] = patch_site(tab, parts[part].decode("utf-8")).encode("utf-8")


# ----------------------------------------------------------------------------- Validation tab
def patch_validation(xml):
    rows = defaultdict(list)
    for m in re.finditer(r'<row r="(\d+)"[^>]*>(.*?)</row>', xml, re.S):
        rows[int(m.group(1))].append(m.group(2))
    add = defaultdict(list)
    for i, v in enumerate(NEW_ACTIONS):
        add[6 + i].append(cell_xml(f"B{6 + i}", v, 49))
    for col, (head, values) in LISTS.items():
        add[1].append(cell_xml(f"{col}1", head, 64))
        for i, v in enumerate(values):
            add[2 + i].append(cell_xml(f"{col}{2 + i}", v, 49))
    out = []
    for r in sorted(set(rows) | set(add)):
        out.append(f'<row r="{r}">{"".join(rows.get(r, []))}{"".join(add.get(r, []))}</row>')
    xml = re.sub(r"<sheetData>.*</sheetData>", "<sheetData>" + "".join(out) + "</sheetData>", xml, flags=re.S)
    cols = "".join(f'<col customWidth="1" min="{col_num(c)}" max="{col_num(c)}" width="{w}"/>'
                   for c, w in (("C", 3), ("D", 36), ("E", 14), ("F", 30), ("G", 10), ("H", 22), ("I", 34)))
    xml = xml.replace("</cols>", cols + "</cols>", 1)
    return xml


parts[VALIDATION_XML] = patch_validation(parts[VALIDATION_XML].decode("utf-8")).encode("utf-8")


# ----------------------------------------------------------------------------- new sheets
class Sheet:
    def __init__(self, name, tab_color, gridlines=False):
        self.name, self.tab_color, self.gridlines = name, tab_color, gridlines
        self.cells = defaultdict(dict)  # row -> col -> (value, style)
        self.widths, self.heights, self.merges = {}, {}, []
        self.cf, self.dv = [], []
        self.freeze = None  # (cols, rows)
        self.autofilter = None
        self.hidden_cols = set()

    def set(self, ref, value, style=None):
        m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
        self.cells[int(m.group(2))][m.group(1)] = (value, X[style] if isinstance(style, str) else style)

    def row(self, r, start_col, values, style=None):
        c = col_num(start_col)
        for i, v in enumerate(values):
            self.set(f"{col_let(c + i)}{r}", v, style)

    def xml(self):
        out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
               f'<sheetPr><tabColor rgb="{self.tab_color}"/></sheetPr><sheetViews><sheetView workbookViewId="0"'
               f'{"" if self.gridlines else " showGridLines=\"0\""}>']
        if self.freeze:
            fc, fr = self.freeze
            tl = f"{col_let(fc + 1)}{fr + 1}"
            pane = "bottomRight" if fc and fr else ("bottomLeft" if fr else "topRight")
            out.append(f'<pane{f" xSplit=\"{fc}\"" if fc else ""}{f" ySplit=\"{fr}\"" if fr else ""} '
                       f'topLeftCell="{tl}" activePane="{pane}" state="frozen"/>'
                       f'<selection pane="{pane}" activeCell="{tl}" sqref="{tl}"/>')
        out.append('</sheetView></sheetViews><sheetFormatPr defaultRowHeight="15"/>')
        if self.widths:
            out.append("<cols>")
            for c in sorted(self.widths, key=col_num):
                hid = ' hidden="1"' if c in self.hidden_cols else ""
                out.append(f'<col min="{col_num(c)}" max="{col_num(c)}" width="{self.widths[c]}" customWidth="1"{hid}/>')
            out.append("</cols>")
        out.append("<sheetData>")
        for r in sorted(set(self.cells) | set(self.heights)):
            ht = f' ht="{self.heights[r]}" customHeight="1"' if r in self.heights else ""
            cells = "".join(cell_xml(f"{c}{r}", v, s) for c, (v, s) in sorted(self.cells[r].items(), key=lambda kv: col_num(kv[0])))
            out.append(f'<row r="{r}"{ht}>{cells}</row>')
        out.append("</sheetData>")
        if self.autofilter:
            out.append(f'<autoFilter ref="{self.autofilter}"/>')
        if self.merges:
            out.append(f'<mergeCells count="{len(self.merges)}">' + "".join(f'<mergeCell ref="{m}"/>' for m in self.merges) + "</mergeCells>")
        out.append(sheet_cf(self.cf, 1))
        if self.dv:
            out.append(f'<dataValidations count="{len(self.dv)}">' + "".join(dv_xml(*d) for d in self.dv) + "</dataValidations>")
        out.append("</worksheet>")
        return "".join(out)


def days():
    d = FIRST_DAY
    while d <= LAST_DAY:
        yield d
        d += dt.timedelta(days=1)


def q(tab):
    return "'" + tab.replace("'", "''") + "'"


SET_ROW = {s[0]: 14 + i for i, s in enumerate(SITES)}  # Settings row of each tab

# ---- Settings
S = Sheet("Settings", "FF666666")
S.set("A1", "Settings", "title")
S.set("A2", "Inputs have a teal background. Formulas in every report tab read these values.", "sub")
S.row(4, "A", ["Parameter", "Value", "Notes"], "h_calc")
S.set("B4", "Value", "h_input")
params = [
    ("Endpoint threshold (pop-up disconnects ÷ calls handled)", 0.05, "input_pct",
     "A day above this needs a same-day technical check and a written cause (Endpoint Check)."),
    ("Repeat offender window (days)", 30, "input_int",
     "Cases per agent are counted over this many days, ending on the report date."),
    ("Repeat offender minimum cases", 2, "input_int",
     "Agents with at least this many cases in the window are repeat offenders."),
    ("Report date (blank = latest dashboard date)", None, "input_date",
     "Date used by Repeat Offenders. Leave blank to follow RAW."),
    ("Report date used", F("IF(B8=\"\",B10,B8)"), "calc_date", "Latest Call Day in RAW unless a date is typed above."),
    ("Latest dashboard date in RAW", F(f"MAX(RAW!$C$2:$C${RAW_LAST_ROW})"), "calc_date", "Last Call Day loaded in RAW."),
    ("Repeat offender window starts", F("B9-B6+1"), "calc_date", "First day counted for repeat offenders."),
]
for i, (label, val, style, note) in enumerate(params):
    r = 5 + i
    S.set(f"A{r}", label, "text_b")
    S.set(f"B{r}", val, style)
    S.set(f"C{r}", note, "text")
S.row(13, "A", ["Language tab", "Team name in summaries", "Cost center", "Cost center in RAW contains",
                "Language in RAW (market)"], "h_input")
S.set("A13", "Language tab", "h_calc")
for tab, _, team, cc, ccm, lang in SITES:
    r = SET_ROW[tab]
    S.set(f"A{r}", tab, "text_b")
    S.row(r, "B", [team, cc, ccm, lang], "input")
S.set("A20", "RAW columns P (Cost Center) and Q (Language) come from the Convex roster. The values above "
             "match the QUERY in cell A1 of each language tab.", "note")
S.merges.append("A20:E20")
S.heights[20] = 32
S.widths.update({"A": 46, "B": 22, "C": 70, "D": 26, "E": 24})
S.dv += [("B5", "decimal", "0", True, "greaterThan"), ("B6:B7", "whole", "0", True, "greaterThan"),
         ("B8", "date", "1", True, "greaterThan")]

# ---- Endpoint Check
EC = Sheet("Endpoint Check", "FFE69138")
EC_FIRST = 5
ec_rows = []
EC.set("A1", "Endpoint check: pop-up disconnects against the threshold", "title")
EC.set("A2", F('"Threshold "&TEXT(Settings!$B$5,"0%")&". Pop-up disconnects come from RAW (Talk Disconnect = '
               'endpoint). Calls handled is not in the extract: type it from the dashboard every day, or type the '
               'dashboard Endpoint % in column F. A day above the threshold needs a technical check on the same day '
               'and a written cause."'), "sub")
EC.row(3, "A", ["Day", "", "", "Dashboard", "Volume", "", "Endpoint %", "", "", "Technical check (same day)",
                "", "", "", "", "Status", ""], "group")
EC.merges += ["A3:C3", "E3:F3", "G3:I3", "J3:N3"]
heads = [("Date", "h_calc"), ("Day", "h_calc"), ("Language tab", "h_calc"), ("Pop-up disconnects", "h_calc"),
         ("Calls handled", "h_input"), ("Endpoint % on dashboard (optional)", "h_input"), ("Endpoint %", "h_calc"),
         ("Threshold", "h_calc"), ("Above threshold?", "h_calc"), ("Technical check done?", "h_input"),
         ("Check date", "h_input"), ("Checked by", "h_input"), ("Cause category", "h_input"),
         ("Written cause", "h_input"), ("Status", "h_calc"), ("Key", "h_calc")]
for i, (h, s) in enumerate(heads):
    EC.set(f"{col_let(i + 1)}4", h, s)
r = EC_FIRST
for d in days():
    for tab, *_ in SITES:
        sr = SET_ROW[tab]
        EC.set(f"A{r}", d, "date")
        EC.set(f"B{r}", F(en_day(f"$A{r}", True)), "calc_c")
        EC.set(f"C{r}", tab, "text")
        EC.set(f"D{r}", F(f'COUNTIFS(RAW!$C$2:$C${RAW_LAST_ROW},$A{r},RAW!$G$2:$G${RAW_LAST_ROW},"endpoint",'
                          f'RAW!$P$2:$P${RAW_LAST_ROW},"*"&Settings!$D${sr}&"*",RAW!$Q$2:$Q${RAW_LAST_ROW},Settings!$E${sr})'),
               "calc_int")
        EC.set(f"E{r}", None, "input_int")
        EC.set(f"F{r}", None, "input_pct")
        EC.set(f"G{r}", F(f'IF(F{r}<>"",F{r},IF(N(E{r})>0,D{r}/E{r},""))'), "calc_pct")
        EC.set(f"H{r}", F("Settings!$B$5"), "calc_pct")
        EC.set(f"I{r}", F(f'IF(G{r}="",IF(D{r}>0,"Volume missing",""),IF(G{r}>H{r},"Yes","No"))'), "calc_c")
        for c, s in (("J", "input_c"), ("K", "input_date"), ("L", "input"), ("M", "input"), ("N", "input_wrap")):
            EC.set(f"{c}{r}", None, s)
        EC.set(f"O{r}", F(f'IF(I{r}="",IF($A{r}>Settings!$B$10,"","No pop-up disconnects"),'
                          f'IF(I{r}="Volume missing","Enter calls handled or the dashboard %",'
                          f'IF(I{r}="No","Below threshold",IF(J{r}<>"Yes","Technical check required",'
                          f'IF(N{r}="","Written cause missing",IF(K{r}="","Check date missing",'
                          f'IF(K{r}=A{r},"Checked same day",IF(K{r}>A{r},"Checked "&(K{r}-A{r})&" day(s) late",'
                          f'"Check date is before the call day"))))))))'), "calc")
        EC.set(f"P{r}", F(f'TEXT($A{r},"yyyy-mm-dd")&"|"&$C{r}'), "key")
        ec_rows.append((d, tab, r))
        r += 1
EC_LAST = r - 1
EC.freeze = (3, 4)
EC.autofilter = f"A4:P{EC_LAST}"
EC.widths.update({"A": 12, "B": 6, "C": 13, "D": 11, "E": 10, "F": 13, "G": 10, "H": 10, "I": 11, "J": 11,
                  "K": 12, "L": 16, "M": 26, "N": 46, "O": 30, "P": 22})
ECR = f"5:{EC_LAST}"
EC.cf += [
    (f"A{EC_FIRST}:C{EC_LAST}", "expr", f"WEEKDAY($A{EC_FIRST},2)>5", DXF["weekend"]),
    (f"I{EC_FIRST}:I{EC_LAST}", "expr", f'$I{EC_FIRST}="Yes"', DXF["red"]),
    (f"O{EC_FIRST}:O{EC_LAST}", "expr", f'$O{EC_FIRST}="Checked same day"', DXF["green"]),
    (f"O{EC_FIRST}:O{EC_LAST}", "expr", f'LEFT($O{EC_FIRST},8)="Checked "', DXF["yellow"]),
    (f"O{EC_FIRST}:O{EC_LAST}", "expr", f'OR($O{EC_FIRST}="Below threshold",$O{EC_FIRST}="No pop-up disconnects")', DXF["grey"]),
    (f"O{EC_FIRST}:O{EC_LAST}", "expr", f'$O{EC_FIRST}<>""', DXF["orange"]),
    (f"D{EC_FIRST}:D{EC_LAST}", "cell", ("equal", "0"), DXF["zero"]),
]
EC.dv += [(f"E{EC_FIRST}:E{EC_LAST}", "whole", "0", True, "greaterThanOrEqual"),
          (f"F{EC_FIRST}:F{EC_LAST}", "decimal", "0", True, "greaterThanOrEqual"),
          (f"J{EC_FIRST}:J{EC_LAST}", "list", lst("G")),
          (f"K{EC_FIRST}:K{EC_LAST}", "date", "1", True, "greaterThan"),
          (f"M{EC_FIRST}:M{EC_LAST}", "list", lst("I"))]

# ---- All Cases
AC = Sheet("All Cases", "FF607D8B")
AC.set("A1", "All cases: every dashboard case from the five language tabs", "title")
AC.set("A2", "Read-only list built by a formula in A4, sorted by Call Day. Edit reviews in the language tabs. "
             "Use Data > Create a filter or a filter view to look at one tab, agent or root cause.", "sub")
AC_COLS = [  # header, source column on the language tab (None = tab name), style
    ("Language tab", None, "text"), ("Id Interaction", "A", "text"), ("Id Case Number", "B", "text"),
    ("Call Day", "C", "date"), ("Case type", "U", "text"), ("Hung up (final)", "G", "text"),
    ("Talk Disconnect", "F", "text"), ("Agent LDAP", "I", "text"), ("Team lead", "J", "text"),
    ("Sector", "K", "text"), ("Validator", "N", "text"), ("Validation date", "O", "date"),
    ("Validation", "P", "text"), ("Root cause", "V", "text"), ("Callback made?", "W", "text"),
    ("No-callback reason", "X", "text"), ("Action taken", "R", "text"), ("Action date", "Y", "date"),
    ("Comment", "Q", "text"), ("Review check", "Z", "text"), ("Agent cases (repeat window)", "AA", "text"),
]
ACL = {h: col_let(i + 1) for i, (h, _, _) in enumerate(AC_COLS)}
for i, (h, _, _) in enumerate(AC_COLS):
    AC.set(f"{col_let(i + 1)}3", h, "h_calc")


def tab_block(tab):
    rng = lambda c: f"{q(tab)}!${c}$2:${c}${SITE_LAST_ROW}"
    pieces = [f'IF({rng("A")}<>"","{tab}","")' if src is None else rng(src) for _, src, _ in AC_COLS]
    return f'IFERROR(FILTER(HSTACK({",".join(pieces)}),{rng("A")}<>""),"")'


ac_formula = (f'LET(v_all,VSTACK({",".join(tab_block(s[0]) for s in SITES)}),'
              f'SORT(FILTER(v_all,IFERROR(INDEX(v_all,0,2)<>"",FALSE)),4,{{ASC}}))')
n_cases = sum(len(v) for v in IDS.values())


def expected_repeat_count():
    """Rows of the Repeat Offenders list (only used to size the array in the LibreOffice test file)."""
    st = (TEST_INPUTS or {}).get("settings", {})
    raw_days = [r[2] for r in wb_vals["RAW"].iter_rows(min_row=2, max_col=3, values_only=True) if isinstance(r[2], dt.datetime)]
    hi = dt.date.fromisoformat(st["B8"]) if st.get("B8") else max(raw_days).date()
    lo = hi - dt.timedelta(days=st.get("B6", 30) - 1)
    n = defaultdict(int)
    for tab, *_ in SITES:
        for r in wb_vals[tab].iter_rows(min_row=2, max_row=SITE_LAST_ROW, max_col=9, values_only=True):
            if r[0] and isinstance(r[2], dt.datetime) and lo <= r[2].date() <= hi and r[8]:
                n[r[8]] += 1
    return sum(v >= st.get("B7", 2) for v in n.values())
AC.set("A4", F(ac_formula, f"A4:{col_let(len(AC_COLS))}{3 + n_cases}"), "text")
for r in range(5, CASES_LAST_ROW + 1):
    for h in ("Call Day", "Validation date", "Action date"):
        AC.set(f"{ACL[h]}{r}", None, "date")
for h in ("Call Day", "Validation date", "Action date"):
    AC.set(f"{ACL[h]}4", None, "date")
AC.set(f"A{CASES_LAST_ROW + 1}", f"End of the formula area (the list can grow to row {CASES_LAST_ROW}).", "end")
AC.freeze = (2, 3)
AC.widths.update({"A": 13, "B": 36, "C": 14, "D": 12, "E": 17, "F": 9, "G": 10, "H": 24, "I": 22, "J": 15,
                  "K": 18, "L": 12, "M": 18, "N": 30, "O": 10, "P": 26, "Q": 14, "R": 12, "S": 50, "T": 28, "U": 11})


def ac(h, absolute=True):
    c = ACL[h]
    return f"'All Cases'!${c}$4:${c}${CASES_LAST_ROW}"


# ---- Process Tracker
PT = Sheet("Process Tracker", "FF8E7CC3")
PT_FIRST, PT_LAST = 5, 4 + PT_ROWS
PT.set("A1", "Process tracker: repeat offenders in a coaching or disciplinary process", "title")
PT.set("A2", "One row per agent, kept up to date by the team lead or OM. Repeat Offenders and the Daily Summary "
             "read In process?, Process stage and Date of last action from here.", "sub")
for i, (h, s) in enumerate([("Agent LDAP", "h_input"), ("Language tab", "h_calc"), ("In process?", "h_input"),
                            ("Process stage", "h_input"), ("Date of last action", "h_input"),
                            ("Owner (TL / OM)", "h_input"), ("Next step / notes", "h_input"),
                            ("Updated on", "h_input"), ("Status shown in summaries", "h_calc")]):
    PT.set(f"{col_let(i + 1)}4", h, s)
for r in range(PT_FIRST, PT_LAST + 1):
    PT.set(f"A{r}", None, "input")
    PT.set(f"B{r}", F(f'IF($A{r}="","",IFERROR(INDEX({ac("Language tab")},MATCH($A{r},{ac("Agent LDAP")},0)),'
                      f'"not in case log"))'), "calc")
    for c, s in (("C", "input_c"), ("D", "input"), ("E", "input_date"), ("F", "input"), ("G", "input_wrap"),
                 ("H", "input_date")):
        PT.set(f"{c}{r}", None, s)
    PT.set(f"I{r}", F(f'IF($A{r}="","",IF(C{r}="Yes",IF(D{r}="","in process",D{r})&IF(E{r}="",""," (last action "'
                      f'&{en_date(f"E{r}")}&")"),IF(C{r}="No","not in process","process not confirmed")))'), "calc")
if TEST_INPUTS:
    for i, rowvals in enumerate(TEST_INPUTS.get("tracker", [])):
        for c, v in rowvals.items():
            if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                v = dt.date.fromisoformat(v)
            PT.set(f"{c}{PT_FIRST + i}", v, PT.cells[PT_FIRST + i][c][1])
PT.freeze = (1, 4)
PT.widths.update({"A": 26, "B": 14, "C": 11, "D": 20, "E": 14, "F": 22, "G": 44, "H": 12, "I": 42})
PT.cf += [(f"A{PT_FIRST}:A{PT_LAST}", "expr", f'AND($A{PT_FIRST}<>"",COUNTIF($A${PT_FIRST}:$A${PT_LAST},$A{PT_FIRST})>1)',
           DXF["red"])]
PT.dv += [(f"A{PT_FIRST}:A{PT_LAST}", "list", f"'Repeat Offenders'!$A$5:$A${4 + RO_ROWS}", False),
          (f"C{PT_FIRST}:C{PT_LAST}", "list", lst("G")), (f"D{PT_FIRST}:D{PT_LAST}", "list", lst("H")),
          (f"E{PT_FIRST}:E{PT_LAST}", "date", "1", True, "greaterThan"),
          (f"H{PT_FIRST}:H{PT_LAST}", "date", "1", True, "greaterThan")]
PTR = lambda c: f"'Process Tracker'!${c}${PT_FIRST}:${c}${PT_LAST}"

# ---- Repeat Offenders
RO = Sheet("Repeat Offenders", "FFCC4125")
RO_FIRST, RO_LAST = 5, 4 + RO_ROWS
RO.set("A1", "Repeat offenders: agents with more than one case", "title")
RO.set("A2", F('"Agents with "&Settings!$B$7&" or more cases from "&' + en_date("Settings!$B$11") + '&" to "&'
               + en_date("Settings!$B$9") + '&" ("&Settings!$B$6&" days), all language tabs. Window and report '
               'date are in Settings."'), "sub")
RO.set("A3", "In process?, Process stage and Date of last action come from Process Tracker. Every agent marked "
             "Confirm process needs a row there.", "sub")
ro_heads = [("Agent LDAP", "h_calc"), ("Cases in window", "h_calc"), ("Language tab", "h_calc"),
            ("Team lead", "h_calc"), ("Hung up (final)", "h_calc"), ("Behaviour (Specialist issues)", "h_calc"),
            ("Not yet reviewed", "h_calc"), ("First case in window", "h_calc"), ("Last case", "h_calc"),
            ("Last action on a case", "h_calc"), ("Last action date on a case", "h_calc"),
            ("In process?", "h_calc"), ("Process stage", "h_calc"), ("Date of last action (tracker)", "h_calc"),
            ("Owner", "h_calc"), ("Date of last action", "h_calc"), ("Status", "h_calc")]
for i, (h, s) in enumerate(ro_heads):
    RO.set(f"{col_let(i + 1)}4", h, s)
win = (f'{ac("Call Day")},">="&Settings!$B$11,{ac("Call Day")},"<="&Settings!$B$9')
RO.set("A5", F(f'IFERROR(LET(v_ag,FILTER({ac("Agent LDAP")},({ac("Call Day")}>=Settings!$B$11)*'
               f'({ac("Call Day")}<=Settings!$B$9)*({ac("Agent LDAP")}<>"")),v_u,UNIQUE(v_ag),'
               f'v_n,COUNTIF(v_ag,v_u),SORT(FILTER(HSTACK(v_u,v_n),v_n>=Settings!$B$7),2,{{DESC}})),"")',
               f"A5:B{RO_LAST if DIALECT == 'google' else 4 + max(1, expected_repeat_count())}"), "text")
for r in range(RO_FIRST, RO_LAST + 1):
    a = f"$A{r}"
    if r > RO_FIRST:
        RO.set(f"A{r}", None, "text")
    RO.set(f"B{r}", None, "calc_int")
    RO.set(f"C{r}", F(f'IF({a}="","",IFERROR(INDEX({ac("Language tab")},MATCH({a},{ac("Agent LDAP")},0)),""))'), "calc")
    RO.set(f"D{r}", F(f'IF({a}="","",IFERROR(INDEX({ac("Team lead")},MATCH({a},{ac("Agent LDAP")},0)),""))'), "calc")
    RO.set(f"E{r}", F(f'IF({a}="","",COUNTIFS({ac("Agent LDAP")},{a},{ac("Hung up (final)")},1,{win}))'), "calc_int")
    RO.set(f"F{r}", F(f'IF({a}="","",COUNTIFS({ac("Agent LDAP")},{a},{ac("Validation")},"Specialist issues",{win}))'),
           "calc_int")
    RO.set(f"G{r}", F(f'IF({a}="","",COUNTIFS({ac("Agent LDAP")},{a},{ac("Validation")},"",{win}))'), "calc_int")
    RO.set(f"H{r}", F(f'IF({a}="","",MINIFS({ac("Call Day")},{ac("Agent LDAP")},{a},{win}))'), "calc_date")
    RO.set(f"I{r}", F(f'IF({a}="","",MAXIFS({ac("Call Day")},{ac("Agent LDAP")},{a},{win}))'), "calc_date")
    RO.set(f"J{r}", F(f'IF(OR({a}="",K{r}=""),"",IFERROR(INDEX(FILTER({ac("Action taken")},({ac("Agent LDAP")}={a})*'
                      f'({ac("Action date")}=K{r})),1),""))', f"J{r}"), "calc")
    RO.set(f"K{r}", F(f'IF({a}="","",IF(MAXIFS({ac("Action date")},{ac("Agent LDAP")},{a})=0,"",'
                      f'MAXIFS({ac("Action date")},{ac("Agent LDAP")},{a})))'), "calc_date")
    RO.set(f"L{r}", F(f'IF({a}="","",IFERROR(INDEX({PTR("C")},MATCH({a},{PTR("A")},0))&"",""))'), "calc_c")
    RO.set(f"M{r}", F(f'IF({a}="","",IFERROR(INDEX({PTR("D")},MATCH({a},{PTR("A")},0))&"",""))'), "calc")
    RO.set(f"N{r}", F(f'IF({a}="","",IF(MAXIFS({PTR("E")},{PTR("A")},{a})=0,"",MAXIFS({PTR("E")},{PTR("A")},{a})))'),
           "calc_date")
    RO.set(f"O{r}", F(f'IF({a}="","",IFERROR(INDEX({PTR("F")},MATCH({a},{PTR("A")},0))&"",""))'), "calc")
    RO.set(f"P{r}", F(f'IF({a}="","",IF(N{r}<>"",N{r},K{r}))'), "calc_date")
    RO.set(f"Q{r}", F(f'IF({a}="","",IF(L{r}="","Confirm process",IF(L{r}="No","Not in process: decide next step",'
                      f'IF(P{r}="","Date of last action missing",IF(I{r}>P{r},"New case after last action",'
                      f'"In process, up to date")))))'), "calc")
RO.set(f"A{RO_LAST + 1}", f"End of the formula area (the list can grow to row {RO_LAST}).", "end")
RO.freeze = (1, 4)
RO.widths.update({"A": 26, "B": 9, "C": 13, "D": 22, "E": 9, "F": 12, "G": 10, "H": 12, "I": 12, "J": 20,
                  "K": 13, "L": 10, "M": 18, "N": 13, "O": 18, "P": 13, "Q": 30})
RO.cf += [(f"Q{RO_FIRST}:Q{RO_LAST}", "expr", f'$Q{RO_FIRST}="In process, up to date"', DXF["green"]),
          (f"Q{RO_FIRST}:Q{RO_LAST}", "expr", f'$Q{RO_FIRST}<>""', DXF["orange"]),
          (f"E{RO_FIRST}:G{RO_LAST}", "cell", ("equal", "0"), DXF["zero"])]

# ---- Daily Summary
DS = Sheet("Daily Summary", "FF3D85C6")
DS_FIRST = 5
DS.set("A1", "Daily summary: every day, every language tab", "title")
DS.set("A2", "One row per day and language tab, weekends included. Every dashboard case is counted: the parts of "
             "each group add up to the cases in the tab. Teal columns are typed; the rest are formulas.", "sub")
groups = [("A3:C3", "Day"), ("D3:F3", "Dashboard reconciliation"), ("G3:I3", "Case type"),
          ("J3:O3", "Review and root cause"), ("P3:S3", "Callback"), ("T3:Y3", "Action taken"),
          ("Z3:AB3", "Details"), ("AC3:AD3", "Endpoint check"), ("AE3:AF3", "Repeat offenders"),
          ("AG3:AI3", "Typed by the team"), ("AJ3:AL3", "Summary")]
for rng, label in groups:
    a, b = rng.split(":")
    for c in range(col_num(a[:-1]), col_num(b[:-1]) + 1):
        DS.set(f"{col_let(c)}3", label if col_let(c) == a[:-1] else None, "group")
    DS.merges.append(rng)
ds_heads = [
    ("A", "Date", "h_calc"), ("B", "Day", "h_calc"), ("C", "Language tab", "h_calc"),
    ("D", "Dashboard cases (RAW)", "h_calc"), ("E", "Cases in tab", "h_calc"), ("F", "Dashboard vs tab", "h_calc"),
    ("G", "Disconnected", "h_calc"), ("H", "Hung up by specialist (final)", "h_calc"),
    ("I", "of which pop-up / endpoint", "h_calc"),
    ("J", "Reviewed", "h_calc"), ("K", "Not yet reviewed", "h_calc"), ("L", "Technical", "h_calc"),
    ("M", "Behaviour", "h_calc"), ("N", "Valid disconnection", "h_calc"), ("O", "Under investigation", "h_calc"),
    ("P", "Callback made", "h_calc"), ("Q", "No callback", "h_calc"), ("R", "Callback N/A", "h_calc"),
    ("S", "Callback not recorded", "h_calc"),
    ("T", "Coached", "h_calc"), ("U", "Removal", "h_calc"), ("V", "IT ticket / feedback", "h_calc"),
    ("W", "Pending", "h_calc"), ("X", "Reviewed, no action taken", "h_calc"), ("Y", "No action recorded", "h_calc"),
    ("Z", "Technical causes", "h_calc"), ("AA", "Behaviour causes", "h_calc"), ("AB", "No-callback reasons", "h_calc"),
    ("AC", "Endpoint %", "h_calc"), ("AD", "Endpoint check status", "h_calc"),
    ("AE", "Repeat offenders with a case this day", "h_calc"), ("AF", "Repeat offenders: process status", "h_calc"),
    ("AG", "Additional context", "h_input"), ("AH", "Prepared by", "h_input"), ("AI", "Sent on", "h_input"),
    ("AJ", "Summary text", "h_calc"), ("AK", "Coverage check", "h_calc"), ("AL", "Key", "h_calc"),
]
for c, h, s in ds_heads:
    DS.set(f"{c}4", h, s)


def site_rng(tab, c):
    return f"{q(tab)}!${c}$2:${c}${SITE_LAST_ROW}"


def textjoin_counts(tab, r, col, values_rng, prefix, total):
    """'2 Headset, 1 Network / connection, 3 not recorded' for the day (array formula)."""
    return F(f'IF($E{r}=0,"",LET(v_n,COUNTIFS({site_rng(tab, "C")},$A{r},{site_rng(tab, col)},{values_rng}),'
             f'TEXTJOIN(", ",TRUE,IF(v_n>0,v_n&" "&SUBSTITUTE({values_rng},"{prefix}",""),""),'
             f'IF({total}-SUM(v_n)>0,{total}-SUM(v_n)&" not recorded",""))))', f"{{cell}}")


def pct(n, base="$E"):
    return f'TEXT(IF({base}=0,0,{n}/{base}),"0%")'


ds_rows = []
r = DS_FIRST
for d in days():
    first_of_day = r
    for tab, _, team, cc, ccm, lang in SITES + [(OTHER, None, None, None, None, None)]:
        R = r
        DS.set(f"A{R}", d, "date")
        DS.set(f"B{R}", F(en_day(f"$A{R}", True)), "calc_c")
        DS.set(f"C{R}", tab, "text_i" if tab == OTHER else "text")
        DS.set(f"AL{R}", F(f'TEXT($A{R},"yyyy-mm-dd")&"|"&$C{R}'), "key")
        for c in ("AG", "AH"):
            DS.set(f"{c}{R}", None, "input_wrap")
        DS.set(f"AI{R}", None, "input_date")
        if tab == OTHER:
            DS.set(f"D{R}", F(f"COUNTIFS(RAW!$C$2:$C${RAW_LAST_ROW},$A{R})-SUM(D{first_of_day}:D{R - 1})"), "calc_int")
            DS.set(f"F{R}", F(f'IF(D{R}=0,"OK",D{R}&" not in any tab")'), "calc")
            for c in [x for x, _, _ in ds_heads[4:]]:
                if c in ("F", "AG", "AH", "AI", "AL"):
                    continue
                DS.set(f"{c}{R}", None, "calc")
            DS.set(f"AJ{R}", F(f'IF(D{R}=0,"","On "&{en_day(f"$A{R}")}&" "&{en_date(f"$A{R}")}&", "&D{R}&" dashboard case(s) belong '
                               f'to agents whose cost center and language have no language tab (for example Sitel Portugal '
                               f'English, Lille, Gdansk or Cape Town). They are not reviewed in this file.")'), "calc_wrap")
            DS.set(f"AK{R}", F(f'IF(D{R}=0,"OK",D{R}&" case(s) without a language tab")'), "calc")
            ds_rows.append((d, tab, R))
            r += 1
            continue
        sr = SET_ROW[tab]
        cnt = lambda *crit: (f"COUNTIFS({site_rng(tab, 'C')},$A{R}" +
                             "".join(f",{site_rng(tab, c)},{v}" for c, v in crit) + ")")
        g0 = lambda *crit: F(f"IF($E{R}=0,0,{cnt(*crit)})")
        DS.set(f"D{R}", F(f'COUNTIFS(RAW!$C$2:$C${RAW_LAST_ROW},$A{R},RAW!$P$2:$P${RAW_LAST_ROW},"*"&Settings!$D${sr}&"*",'
                          f'RAW!$Q$2:$Q${RAW_LAST_ROW},Settings!$E${sr})'), "calc_int")
        DS.set(f"E{R}", F(cnt()), "calc_int")
        DS.set(f"F{R}", F(f'IF(D{R}=E{R},"OK",IF(D{R}>E{R},D{R}-E{R}&" missing in tab",E{R}-D{R}&" extra in tab"))'), "calc")
        DS.set(f"G{R}", g0(("G", "0")), "calc_int")
        DS.set(f"H{R}", g0(("G", "1")), "calc_int")
        DS.set(f"I{R}", g0(("F", '"endpoint"')), "calc_int")
        DS.set(f"J{R}", g0(("P", '"<>"')), "calc_int")
        DS.set(f"K{R}", F(f"E{R}-J{R}"), "calc_int")
        DS.set(f"L{R}", g0(("P", '"Technical issues"')), "calc_int")
        DS.set(f"M{R}", g0(("P", '"Specialist issues"')), "calc_int")
        DS.set(f"N{R}", g0(("P", '"False Flag (Valid Disconnection)"')), "calc_int")
        DS.set(f"O{R}", g0(("P", '"Under Investigation"')), "calc_int")
        DS.set(f"P{R}", g0(("W", '"Yes"')), "calc_int")
        DS.set(f"Q{R}", g0(("W", '"No"')), "calc_int")
        DS.set(f"R{R}", g0(("W", '"N/A"')), "calc_int")
        DS.set(f"S{R}", F(f"E{R}-P{R}-Q{R}-R{R}"), "calc_int")
        DS.set(f"T{R}", g0(("R", '"Coached"')), "calc_int")
        DS.set(f"U{R}", g0(("R", '"Removal"')), "calc_int")
        DS.set(f"V{R}", F(f"IF($E{R}=0,0,{cnt(('R', 'Validation!$B$6'))}+{cnt(('R', 'Validation!$B$7'))})"), "calc_int")
        DS.set(f"W{R}", g0(("R", '"Pending"')), "calc_int")
        DS.set(f"X{R}", g0(("R", '"N/A"')), "calc_int")
        DS.set(f"Y{R}", F(f"E{R}-T{R}-U{R}-V{R}-W{R}-X{R}"), "calc_int")
        for c, col, rng, pre, tot in (("Z", "V", TECH_RANGE, "Technical - ", f"$L{R}"),
                                      ("AA", "V", BEH_RANGE, "Behaviour - ", f"$M{R}"),
                                      ("AB", "X", lst("F"), "", f"$Q{R}")):
            f = textjoin_counts(tab, R, col, rng, pre, tot)
            f.array = f"{c}{R}"
            DS.set(f"{c}{R}", f, "calc_wrap")
        DS.set(f"AC{R}", F(f"IFERROR(INDEX('Endpoint Check'!$G:$G,MATCH($AL{R},'Endpoint Check'!$P:$P,0)),\"\")"), "calc_pct")
        DS.set(f"AD{R}", F(f"IFERROR(INDEX('Endpoint Check'!$O:$O,MATCH($AL{R},'Endpoint Check'!$P:$P,0))&\"\",\"\")"), "calc")
        rep = (f"FILTER({site_rng(tab, 'I')},({site_rng(tab, 'C')}=$A{R})*"
               f"(IFERROR(--{site_rng(tab, 'AA')},0)>=Settings!$B$7))")
        DS.set(f"AE{R}", F(f"IF($E{R}=0,0,IFERROR(ROWS(UNIQUE({rep})),0))", f"AE{R}"), "calc_int")
        DS.set(f"AF{R}", F(f'IF(AE{R}=0,"",LET(v_u,UNIQUE({rep}),TEXTJOIN("; ",TRUE,v_u&" ("&'
                           f'XLOOKUP(v_u,{PTR("A")},{PTR("I")},"process not confirmed")&")")))', f"AF{R}"), "calc_wrap")
        p = lambda n: pct(n, f"$E{R}")
        day_txt = f'{en_day(f"$A{R}")}&" "&{en_date(f"$A{R}")}'
        summary = (
            f'IF(AND($D{R}=0,$E{R}=0),IF($A{R}>Settings!$B$10,"No dashboard data loaded yet for "&{day_txt}&".",'
            f'"On "&{day_txt}&", the "&Settings!$E${sr}&" team ("&Settings!$B${sr}&") recorded no '
            f'disconnected calls."),'
            f'"On "&{day_txt}&", the "&Settings!$E${sr}&" team ("&Settings!$B${sr}&") recorded a '
            f'total of "&$D{R}&" disconnected calls on the dashboard ("&$I{R}&" pop-up / endpoint); "&$H{R}&" of them '
            f'were hung up by the specialist (final)."'
            f'&IF($D{R}<>$E{R},CHAR(10)&"• Dashboard check: "&$E{R}&" of "&$D{R}&" cases are in the {tab} tab.","")'
            f'&CHAR(10)&"• Technical Issues: "&$L{R}&" ("&{p(f"$L{R}")}&")"&IF($Z{R}<>""," - "&$Z{R},"")'
            f'&CHAR(10)&"• Ambassador Behaviours: "&$M{R}&" ("&{p(f"$M{R}")}&")"&IF($AA{R}<>""," - "&$AA{R},"")'
            f'&IF($K{R}+$N{R}+$O{R}>0,CHAR(10)&"• Other cases: "&TEXTJOIN(", ",TRUE,IF($N{R}>0,$N{R}&" valid '
            f'disconnection",""),IF($O{R}>0,$O{R}&" under investigation",""),IF($K{R}>0,$K{R}&" not yet reviewed","")),"")'
            f'&CHAR(10)&"• Callback Performed: "&$P{R}&" ("&{p(f"$P{R}")}&")"'
            f'&CHAR(10)&"• No Callback Performed: "&$Q{R}&" ("&{p(f"$Q{R}")}&")"&IF($AB{R}<>""," - "&$AB{R},"")'
            f'&IF($R{R}+$S{R}>0,CHAR(10)&"• Callback not applicable: "&$R{R}&"; callback not recorded: "&$S{R},"")'
            f'&CHAR(10)&"• Coachings Delivered: "&IF($T{R}=0,"N/A",$T{R})'
            f'&CHAR(10)&"• Removals from Campaign: "&IF($U{R}=0,"N/A",$U{R})'
            f'&CHAR(10)&"• Other actions: "&$V{R}&" IT ticket or feedback, "&$W{R}&" pending, "&$X{R}&" reviewed with '
            f'no action taken"&IF($Y{R}>0,", "&$Y{R}&" with no action recorded yet","")'
            f'&CHAR(10)&"• Pop-up disconnects (threshold "&TEXT(Settings!$B$5,"0%")&"): "&IF($AC{R}="",'
            f'"calls handled not entered",TEXT($AC{R},"0.0%")&" - "&$AD{R}'
            f'&IFERROR(IF(INDEX(\'Endpoint Check\'!$N:$N,MATCH($AL{R},\'Endpoint Check\'!$P:$P,0))="",""," - cause: "'
            f'&INDEX(\'Endpoint Check\'!$N:$N,MATCH($AL{R},\'Endpoint Check\'!$P:$P,0))),""))'
            f'&CHAR(10)&"• Repeat offenders: "&IF($AE{R}=0,"none",$AE{R}&" - "&$AF{R})'
            f'&CHAR(10)&"• Additional Context: "&IF($AG{R}="","N/A",$AG{R}))'
        )
        DS.set(f"AJ{R}", F(summary), "calc_wrap")
        open_checks = "(" + "+".join(cnt(("Z", f'"{z}"')) for z in (
            "Root cause differs from Validation", "No-callback reason missing", "Action date missing")) + ")"
        DS.set(f"AK{R}", F(
            f'IF(AND($D{R}=0,$E{R}=0),IF($A{R}>Settings!$B$10,"No dashboard data yet","No cases"),'
            f'LET(v_t,TEXTJOIN("; ",TRUE,IF($D{R}<>$E{R},"dashboard "&$D{R}&" vs tab "&$E{R},""),'
            f'IF($K{R}>0,$K{R}&" not reviewed",""),'
            f'IF({cnt(("P", chr(34) + "<>" + chr(34)), ("V", chr(34) * 2))}>0,{cnt(("P", chr(34) + "<>" + chr(34)), ("V", chr(34) * 2))}&" reviewed without root cause",""),'
            f'IF($S{R}>0,$S{R}&" without callback info",""),IF($Y{R}>0,$Y{R}&" without action",""),'
            f'IF({open_checks}>0,{open_checks}&" with Review check open",""),'
            f'IF($J{R}<>$L{R}+$M{R}+$N{R}+$O{R},$J{R}-$L{R}-$M{R}-$N{R}-$O{R}&" with a Validation value outside the list","")),'
            f'IF(v_t="","OK: all "&$E{R}&" cases covered",v_t)))'), "calc_wrap")
        ds_rows.append((d, tab, R))
        r += 1
DS_LAST = r - 1
DS.freeze = (3, 4)
DS.autofilter = f"A4:AL{DS_LAST}"
DS.widths.update({"A": 12, "B": 6, "C": 13})
for c in [x for x, _, _ in ds_heads[3:]]:
    DS.widths[c] = 9
DS.widths.update({"F": 14, "Z": 26, "AA": 22, "AB": 24, "AC": 9, "AD": 22, "AE": 11, "AF": 30, "AG": 30, "AH": 14,
                  "AI": 12, "AJ": 70, "AK": 34, "AL": 22})
DS.heights[4] = 45
DS.cf += [
    (f"A{DS_FIRST}:C{DS_LAST}", "expr", f"WEEKDAY($A{DS_FIRST},2)>5", DXF["weekend"]),
    (f"F{DS_FIRST}:F{DS_LAST}", "expr", f'$F{DS_FIRST}="OK"', DXF["green"]),
    (f"F{DS_FIRST}:F{DS_LAST}", "expr", f'$F{DS_FIRST}<>""', DXF["red"]),
    (f"AK{DS_FIRST}:AK{DS_LAST}", "expr", f'LEFT($AK{DS_FIRST},2)="OK"', DXF["green"]),
    (f"AK{DS_FIRST}:AK{DS_LAST}", "expr", f'OR($AK{DS_FIRST}="No cases",$AK{DS_FIRST}="No dashboard data yet")', DXF["grey"]),
    (f"AK{DS_FIRST}:AK{DS_LAST}", "expr", f'$AK{DS_FIRST}<>""', DXF["orange"]),
    (f"AD{DS_FIRST}:AD{DS_LAST}", "expr", f'$AD{DS_FIRST}="Checked same day"', DXF["green"]),
    (f"AD{DS_FIRST}:AD{DS_LAST}", "expr",
     f'OR($AD{DS_FIRST}="Technical check required",$AD{DS_FIRST}="Written cause missing",$AD{DS_FIRST}="Check date missing")',
     DXF["red"]),
    (f"D{DS_FIRST}:Y{DS_LAST}", "cell", ("equal", "0"), DXF["zero"]),
]
DS.dv += [(f"AI{DS_FIRST}:AI{DS_LAST}", "date", "1", True, "greaterThan")]
if TEST_INPUTS:
    by_key = {(d.isoformat(), t): rr for d, t, rr in ds_rows}
    for k, vals in TEST_INPUTS.get("daily", {}).items():
        d, t = k.split("|")
        for c, v in vals.items():
            if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                v = dt.date.fromisoformat(v)
            DS.set(f"{c}{by_key[(d, t)]}", v, DS.cells[by_key[(d, t)]][c][1])
    ec_key = {(d.isoformat(), t): rr for d, t, rr in ec_rows}
    for k, vals in TEST_INPUTS.get("endpoint", {}).items():
        d, t = k.split("|")
        for c, v in vals.items():
            if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                v = dt.date.fromisoformat(v)
            EC.set(f"{c}{ec_key[(d, t)]}", v, EC.cells[ec_key[(d, t)]][c][1])
    for ref, v in TEST_INPUTS.get("settings", {}).items():
        if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            v = dt.date.fromisoformat(v)
        S.set(ref, v, S.cells[int(ref[1:])][ref[0]][1])

# ---- Daily Report
DR = Sheet("Daily Report", "FF0B5394")
DR.set("A1", "Daily report", "title")
DR.set("A2", "Type a date in C3 (blank = latest dashboard date). Copy each text in section 2 into the daily email. "
             "Weekends are reported like any other day.", "sub")
DR.set("B3", "Report date", "text_b")
DR.set("C3", None, "input_date")
DR.set("D3", "Date used", "text_b")
DR.set("E3", F('IF(C3="",Settings!$B$10,C3)'), "calc_date")
DR.set("F3", F(en_day("E3")), "calc")
DR.set("A5", "1. Numbers for the day", "section")
dr_cols = [("Language tab", None), ("Dashboard cases", "D"), ("Cases in tab", "E"), ("Disconnected", "G"),
           ("Hung up (final)", "H"), ("Pop-up / endpoint", "I"), ("Not yet reviewed", "K"), ("Technical", "L"),
           ("Behaviour", "M"), ("Valid", "N"), ("Under investigation", "O"), ("Callback made", "P"),
           ("No callback", "Q"), ("Coached", "T"), ("Removal", "U"), ("Reviewed, no action taken", "X"),
           ("No action recorded", "Y"), ("Endpoint %", "AC"), ("Endpoint check", "AD"), ("Repeat offenders", "AE"),
           ("Coverage check", "AK"), ("Summary sent on", "AI")]
for i, (h, _) in enumerate(dr_cols):
    DR.set(f"{col_let(i + 2)}6", h, "h_calc")
KEY = lambda r: f'TEXT($E$3,"yyyy-mm-dd")&"|"&$B{r}'
tabs_all = [s[0] for s in SITES] + [OTHER]
for i, tab in enumerate(tabs_all):
    r = 7 + i
    DR.set(f"B{r}", tab, "text_b" if tab != OTHER else "text_i")
    for j, (h, src) in enumerate(dr_cols[1:]):
        c = col_let(j + 3)
        style = {"AC": "calc_pct", "AD": "calc", "AK": "calc_wrap", "AI": "calc_date"}.get(src, "calc_int")
        if tab == OTHER and src not in ("D", "AK"):
            DR.set(f"{c}{r}", None, "calc")
            continue
        tail = '&""' if src in ("AD", "AK") else ""
        look = f"INDEX('Daily Summary'!${src}:${src},MATCH({KEY(r)},'Daily Summary'!$AL:$AL,0))"
        if src == "AI":  # a blank date must stay blank (not 0)
            look = f'IF({look}="","",{look})'
        DR.set(f"{c}{r}", F(f"IFERROR({look}{tail},\"\")"), style)
tr = 7 + len(tabs_all)
DR.set(f"B{tr}", "Total", "calc_bold")
for j, (h, src) in enumerate(dr_cols[1:]):
    c = col_let(j + 3)
    if src in ("AC", "AD", "AK", "AI"):
        DR.set(f"{c}{tr}", None, "calc_bold")
    else:
        DR.set(f"{c}{tr}", F(f"SUM({c}7:{c}{tr - 1})"), ST.xf(fill=GREY_LIGHT, bold=True, num=INT, h="center"))
DR.set(f"A{tr + 2}", "2. Summary texts", "section")
for i, tab in enumerate(tabs_all):
    r = tr + 3 + i
    DR.set(f"B{r}", tab, "text_b" if tab != OTHER else "text_i")
    DR.set(f"C{r}", F(f"IFERROR(INDEX('Daily Summary'!$AJ:$AJ,MATCH({KEY(r)},'Daily Summary'!$AL:$AL,0))&\"\",\"\")"),
           "calc_wrap")
    for c in range(4, len(dr_cols) + 2):
        DR.set(f"{col_let(c)}{r}", None, "calc_wrap")
    DR.merges.append(f"C{r}:{col_let(len(dr_cols) + 1)}{r}")
    DR.heights[r] = 230 if tab != OTHER else 40
s7 = tr + 3 + len(tabs_all) + 1
DR.set(f"A{s7}", "3. Summaries sent in the last 7 days (Sent on in Daily Summary)", "section")
DR.set(f"B{s7 + 1}", "Language tab", "h_calc")
for k in range(7):
    c = col_let(3 + k)
    DR.set(f"{c}{s7 + 1}", F(f"$E$3-{6 - k}"), ST.xf(bold=True, color="FFFFFFFF", fill=SLATE, h="center", num=ST.numfmt("ddd d mmm")))
for i, tab in enumerate([s[0] for s in SITES]):
    r = s7 + 2 + i
    DR.set(f"B{r}", tab, "text_b")
    for k in range(7):
        c = col_let(3 + k)
        key = f'TEXT({c}${s7 + 1},"yyyy-mm-dd")&"|"&$B{r}'
        DR.set(f"{c}{r}", F(f"IFERROR(IF(INDEX('Daily Summary'!$AI:$AI,MATCH({key},'Daily Summary'!$AL:$AL,0))<>\"\","
                            f"\"Sent\",IF({c}${s7 + 1}>Settings!$B$10,\"No data yet\",\"Missing\")),\"\")"), "calc_c")
DR.cf += [
    (f"C{s7 + 2}:I{s7 + 6}", "expr", f'C{s7 + 2}="Sent"', DXF["green"]),
    (f"C{s7 + 2}:I{s7 + 6}", "expr", f'C{s7 + 2}="Missing"', DXF["red"]),
    (f"V7:V{tr - 1}", "expr", 'LEFT($V7,2)="OK"', DXF["green"]),
    (f"V7:V{tr - 1}", "expr", '$V7<>""', DXF["orange"]),
    (f"D7:U{tr - 1}", "cell", ("equal", "0"), DXF["zero"]),
]
DR.dv += [("C3", "date", "1", True, "greaterThan")]
DR.widths.update({"A": 3, "B": 14})
for j in range(len(dr_cols) - 1):
    DR.widths[col_let(j + 3)] = 10
DR.widths.update({"T": 22, "V": 34, "W": 12})
DR.heights[6] = 45
DR.freeze = (0, 3)

# ---- Read Me
RM = Sheet("Read Me", "FF274E13")
RM.widths.update({"A": 3, "B": 44, "C": 62, "D": 40})
rm = [
    ("title", "Disconnected calls: daily reporting"),
    ("sub", "Every disconnected call and hang-up on the dashboard is reported every day, including Saturday and "
            "Sunday, for the five language tabs: São Paulo, Portugal PT, Portugal SP, Seville and Milan."),
    None,
    ("section", "What the client asked for and where it is"),
    ("head", ["Request", "Where", "Filled by"]),
    ("row", ["A summary every day, including weekends, for every language tab",
             "Daily Report: pick the date and copy one text per tab. Daily Summary keeps one row per day and tab.",
             "Formulas. Prepared by and Sent on are typed in Daily Summary."]),
    ("row", ["Root cause, callback and actions for every case",
             "Language tabs, columns U to AA, next to Validation, Comment and Action Taken. All Cases lists every case "
             "of the five tabs in one place.", "Reviewer"]),
    ("row", ["Numbers that match the dashboard",
             "Daily Summary columns D to F compare RAW with each tab. Every case is counted: not yet reviewed, "
             "reviewed with no action taken, and so on. Coverage check lists what is missing.", "Formulas"]),
    ("row", ["Root cause on pop-up disconnects above the 5% Endpoint threshold",
             "Endpoint Check: calls handled, Endpoint %, technical check on the same day and written cause.",
             "Site or technical lead"]),
    ("row", ["Repeat offenders: in the process, and date of last action",
             "Repeat Offenders lists agents with 2 or more cases in the window. Process Tracker holds the process status.",
             "Team lead or OM"]),
    None,
    ("section", "Daily routine (every day, Saturday and Sunday included)"),
    ("row", ["1. Load the dashboard data", "Follow HowToExtract. Paste the new rows at the bottom of RAW.", ""]),
    ("row", ["2. Review every new case in the language tabs",
             "Fill Validator, Validation Date, Validation, Comment and Action Taken as before, then Root cause, "
             "Callback made?, No-callback reason (when no callback was made) and Action date (when an action was "
             "taken). Review check (column Z) shows what is missing and should read Complete.", "Reviewer"]),
    ("row", ["3. Check pop-up disconnects",
             "During the day, watch the Endpoint % on the live dashboard: above the threshold, run the technical check "
             "the same day. In Endpoint Check type Calls handled for each tab (or the dashboard Endpoint %), then "
             "Technical check done?, Check date, Checked by, Cause category and Written cause. The written cause "
             "appears in the daily text.", "Site or technical lead"]),
    ("row", ["4. Confirm repeat offenders",
             "In Repeat Offenders, every agent with status Confirm process needs a row in Process Tracker: In process?, "
             "Process stage, Date of last action and Owner.", "Team lead or OM"]),
    ("row", ["5. Send the summary",
             "In Daily Report choose the date, check that Coverage check reads OK for every tab, copy the texts into "
             "the email, then type Prepared by and Sent on in Daily Summary. Notes for the client go in Additional "
             "context (Daily Summary). The row Other (no tab) is for information.",
             "Owner of the day, weekend included (name a weekend owner in Extraction POCs)"]),
    None,
    ("section", "Definitions"),
    ("row", ["Dashboard cases", "Rows in RAW for the day whose cost center and language match the tab (Settings). "
             "Rows of other cost centers or languages are counted on the row Other (no tab).", ""]),
    ("row", ["Disconnected / Hung up (final)", "Hung Up By Specialist (Final) = 0 or 1. The two add up to the "
             "dashboard cases. Pop-up / endpoint: Talk Disconnect = endpoint.", ""]),
    ("row", ["Reviewed", "Validation is filled. Reviewed, no action taken: Action Taken = N/A. No action recorded: "
             "Action Taken is empty.", ""]),
    ("row", ["Root cause", "Technical (headset, microphone, unmute, transfer, network, Genesys, PC), Behaviour, "
             "Valid disconnection or Under investigation. It must agree with Validation: Technical issues = Technical, "
             "Specialist issues = Behaviour, False Flag = Valid.", ""]),
    ("row", ["Endpoint %", "Pop-up disconnects ÷ calls handled, per day and tab. The threshold is in Settings.", ""]),
    ("row", ["Repeat offender", "An agent with at least the minimum number of cases (Settings) in the window "
             "ending on the report date, across all language tabs.", ""]),
    None,
    ("section", "Keeping the numbers right"),
    ("row", ["Review columns are typed next to RAW rows",
             "The language tabs take their rows from RAW with a QUERY. Add new rows at the bottom of RAW and do not "
             "sort or delete rows in RAW, otherwise reviews move to other cases.", ""]),
    ("row", ["RAW rows read by the language tabs", "The QUERY in cell A1 of each language tab reads RAW rows 1 to "
             "6655, the same rows as the RAW lookups (P:R). Daily Summary column F shows a difference as soon as "
             "dashboard cases are missing from a tab.", ""]),
    ("row", ["Pre-filled values", "PREFILL_NOTE", ""]),
    ("row", ["Colours", "Teal headers and cells are typed. Grey headers are formulas: do not type over them.", ""]),
    ("row", ["Adding days", "Daily Summary and Endpoint Check run to 30 Jun 2027. To add days, copy the last day's "
             "rows (six in Daily Summary, five in Endpoint Check) below and change the dates.", ""]),
    ("row", ["Older summary tabs", "Site Summary, PTDaily, ESDaily and ITDaily are unchanged. They count only cases "
             "validated as Technical issues or Specialist issues, so their totals can be lower than the dashboard.", ""]),
]
total_prefill = sum(PREFILL_COUNT.values())
prefill_note = (f"Callback made? was filled for {total_prefill} cases from the earlier review logs (PT_São Paulo, "
                f"PT_Portugal, SP_Seville & Portugal, IT_Milan) where the same Id Interaction was found "
                f"({', '.join(f'{t} {n}' for t, n in PREFILL_COUNT.items())}). For the Milan cases in IT_Milan, Root "
                f"cause and No-callback reason were filled too. Pre-filled cells are in blue italics: the reviewer "
                f"confirms or corrects them during the review.")
r = 2
for item in rm:
    if item is None:
        r += 1
        continue
    kind, val = item
    if kind in ("title", "sub", "section"):
        RM.set(f"B{r}", val, kind)
        if kind == "sub":
            RM.merges.append(f"B{r}:D{r}")
            RM.heights[r] = 32
            RM.set(f"B{r}", val, "note")
    elif kind == "head":
        RM.row(r, "B", val, "h_calc")
    else:
        val = [prefill_note if v == "PREFILL_NOTE" else v for v in val]
        RM.row(r, "B", val, "wrap")
        RM.set(f"B{r}", val[0], ST.xf(bold=True, wrap=True, v="top"))
        RM.heights[r] = max(32, 16 * (1 + max(len(str(v)) // 70 for v in val)))
    r += 1

NEW_SHEETS = [RM, DR, DS, EC, AC, RO, PT, S]

# ----------------------------------------------------------------------------- workbook parts
wbx = parts["xl/workbook.xml"].decode("utf-8")
rels = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
ctypes = parts["[Content_Types].xml"].decode("utf-8")
sheet_ids = [int(i) for i in re.findall(r'<sheet [^>]*sheetId="(\d+)"', wbx)]
rid_nums = [int(i) for i in re.findall(r'Id="rId(\d+)"', rels)]
entries, sid, rid = [], max(sheet_ids), max(rid_nums)
for sh in NEW_SHEETS:
    sid += 1
    rid += 1
    path = f"worksheets/sheet{sid}.xml"
    parts[f"xl/{path}"] = sh.xml().encode("utf-8")
    entries.append(f'<sheet state="visible" name="{esc(sh.name)}" sheetId="{sid}" r:id="rId{rid}"/>')
    rels = rels.replace("</Relationships>", f'<Relationship Id="rId{rid}" Type="http://schemas.openxmlformats.org/'
                                            f'officeDocument/2006/relationships/worksheet" Target="{path}"/></Relationships>')
    ctypes = ctypes.replace("</Types>", f'<Override ContentType="application/vnd.openxmlformats-officedocument.'
                                        f'spreadsheetml.worksheet+xml" PartName="/xl/{path}"/></Types>')
    sh.part = path
# new tabs go right after the last existing visible tab (Milan), before the hidden tabs
anchor = re.search(r'<sheet [^>]*name="Milan"[^>]*/>', wbx)
wbx = wbx[:anchor.end()] + "".join(entries) + wbx[anchor.end():]
names = re.findall(r'<sheet [^>]*name="([^"]+)"', wbx)
pos = {html.unescape(n): i for i, n in enumerate(names)}
assert pos["Portugal SP"] == 12  # defined names below use localSheetId 12
dn = "".join(f'<definedName hidden="1" localSheetId="{pos[sh.name]}" name="_xlnm._FilterDatabase">'
             f"{esc(q(sh.name))}!{re.sub(r'([A-Z]+)(\d+)', r'$\1$\2', sh.autofilter)}</definedName>"
             for sh in NEW_SHEETS if sh.autofilter)
wbx = wbx.replace("<definedNames>", "<definedNames>" + dn, 1)
wbx = wbx.replace("'Portugal SP'!$A$1:$T$2728</definedName>", "'Portugal SP'!$A$1:$AA$2728</definedName>", 1)
parts["xl/workbook.xml"] = wbx.encode("utf-8")
parts["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
parts["[Content_Types].xml"] = ctypes.encode("utf-8")
parts["xl/styles.xml"] = ST.render().encode("utf-8")

with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
    for n in zin.namelist():
        z.writestr(n, parts.pop(n))
    for n, b in parts.items():
        z.writestr(n, b)

summary = {
    "dialect": DIALECT, "prefill": PREFILL_COUNT, "cases": {t: len(v) for t, v in IDS.items()},
    "daily_summary_rows": [DS_FIRST, DS_LAST], "endpoint_rows": [EC_FIRST, EC_LAST],
    "sheets": [sh.name for sh in NEW_SHEETS], "conflicts": len(conflict),
}
print(json.dumps(summary, ensure_ascii=False))
